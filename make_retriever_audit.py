#!/usr/bin/env python3
"""Human re-scoring of retriever rankings. Per language: 40 random questions from the first 200 MKQA items;
top-3 passages from BM25, LaBSE, mE5-large are pooled (deduplicated), shuffled, blinded. Writes sheets + a private key."""
import os
os.environ.setdefault('KMP_DUPLICATE_LIB_OK','TRUE'); os.environ.setdefault('OMP_NUM_THREADS','1'); os.environ.setdefault('TOKENIZERS_PARALLELISM','false')
from sentence_transformers import SentenceTransformer
import json, random, re, sys, gzip, numpy as np, openpyxl, faiss
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0, ".")
from retrieval.bm25_retriever import BM25Retriever
NQ, K = 40, 3
rnd = random.Random(2026); key = {}; sheets = {}
for L in ("ar", "ms"):
    S = json.load(open(f"results/table9_matched_retrieved_{L}.json", encoding="utf-8"))["samples"]; qi = sorted(rnd.sample(range(200), NQ))
    corpus = json.load(open(f"indexes/corpus_{L}.json", encoding="utf-8"))
    bm = BM25Retriever(L); bm.index_corpus(corpus)
    run = {"bm25": {}, "labse": {}, "me5large": {}}
    for i in qi:
        run["bm25"][i] = [p["text"] for p, _ in bm.retrieve(S[i]["query"], K)]
    import os
    for name, model, prefix in (("labse", "sentence-transformers/LaBSE", ""), ("me5large", "intfloat/multilingual-e5-large", "query: ")):
        if name == "labse" and os.environ.get("SKIP_LABSE"): run["labse"] = {i: [] for i in qi}; continue
        fi = faiss.read_index(f"indexes/{name}_{L}.faiss"); M = fi.reconstruct_n(0, fi.ntotal); del fi; pas = json.load(open(f"indexes/{name}_{L}_passages.json", encoding="utf-8"))
        enc = SentenceTransformer(model, device="cpu")
        for i in qi:
            v = enc.encode([prefix + S[i]["query"]], normalize_embeddings=True)
            sc = M @ np.asarray(v, dtype="float32")[0]; ids = np.argsort(-sc)[:K]; run[name][i] = [pas[int(j)]["text"] for j in ids]
        print(L, name, "done", flush=True)
    rows = []
    for i in qi:
        pool = {}
        for name in run:
            for r, t in enumerate(run[name][i], 1): pool.setdefault(t, {})[name] = r
        for t, ranks in pool.items(): rows.append((i, S[i]["query"], S[i]["reference"], t, ranks))
    rnd.shuffle(rows); sheets[L] = rows
    key[L] = [dict(item=n, q=r[0], ranks=r[4]) for n, r in enumerate(rows, 1)]
    print(L, "pooled pairs:", len(rows), flush=True)
json.dump(key, open("audit/retriever_audit_KEY_do_not_share.json", "w"))
HDR = ["ID", "Question", "Gold answer (MKQA)", "Passage", "Judgment (pick from list)", "Notes (optional)"]
def build(path, subset):
    wb = openpyxl.Workbook(); wb.remove(wb.active); ws0 = wb.create_sheet("Instructions")
    for t in open("audit/INSTRUCTIONS.txt", encoding="utf-8").read().split("\n"): ws0.append([t])
    ws0.append([""]); ws0.append(["TIGHTENED RULE: SUPPORTS only if, reading ONLY the passage, you could answer THIS question and obtain the gold answer. Same number/year/word in an unrelated context = MENTIONS_ONLY."])
    ws0.column_dimensions["A"].width = 130
    for L, rows in sheets.items():
        ws = wb.create_sheet(L.upper()); ws.append(HDR)
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
        dv = DataValidation(type="list", formula1='"SUPPORTS,MENTIONS_ONLY,NO"', allow_blank=True); ws.add_data_validation(dv)
        for n, r in enumerate(rows, 1):
            if subset is not None and n not in subset[L]: continue
            t = r[3]; t = "'" + t if t[:1] in "=+-@" else t
            ws.append([f"{L}-{n:03d}", r[1], r[2], t, None, None]); dv.add(ws.cell(ws.max_row, 5))
        for col, w in zip("ABCDEF", (10, 45, 25, 100, 22, 30)): ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    wb.save(path)
build("../audit_for_raters/retriever_audit_rater1.xlsx", None)
sub = {L: set(random.Random(5).sample(range(1, len(sheets[L]) + 1), max(1, len(sheets[L]) // 4))) for L in sheets}
build("../audit_for_raters/retriever_audit_rater2.xlsx", sub)
json.dump({L: sorted(v) for L, v in sub.items()}, open("audit/retriever_audit_rater2_subset.json", "w"))
print("OK")
